"""
Génère le classeur Excel des prompts ChatGPT pour une ou plusieurs marionnettes 2D (26 images
chacune : base, 9 bouches, 2 clignements, 4 émotions, 3 regards, 3 bouches souriantes, 4 mains).

    python3 scripts/outils/prompts-marionnette.py sortie.xlsx prefixe1 "description 1" [prefixe2 "description 2" …]

Le préfixe est ajouté devant chaque nom de fichier (ex. blazer_base.png) et donne le nom du
dossier assets/<prefixe>/. La description décrit ce qui doit rester identique (vêtements,
coiffure, accessoires) et est insérée dans chaque prompt. Dépendance : openpyxl.
"""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

POST = " Tout le reste doit rester strictement identique, pixel pour pixel si possible : {same}, fond. Image carrée 2048 × 2048, PNG."


def rows_for(prefix: str, desc: str, same: str, hands: bool = True):
    p = prefix + "_"
    pre0 = f"Voici mon personnage de référence (image jointe). Génère exactement la même image — même personnage ({desc}), même cadrage (buste centré, tête au même endroit), même lumière, même fond uni, même style 3D cartoon — avec une seule différence : "
    pre = f"Voici l'image de base de mon personnage (image jointe : {p}base.png). Génère exactement la même image — même personnage ({desc}), même cadrage, même tête, même lumière, même fond uni, même style 3D cartoon — avec une seule différence : "
    post = POST.format(same=same)
    r = [
        (0, f"{p}base.png", "Base", "Neutre : bouche fermée, yeux ouverts droits",
         pre0 + "la bouche est fermée, détendue et neutre (pas de sourire, pas de dents visibles), les yeux ouverts (blanc de l'œil et pupille visibles) regardant droit vers la caméra, sourcils au repos, expression calme et neutre." + post,
         "C'est LA référence : toutes les autres images en dérivent et lui seront comparées. Bouche fermée neutre, regard droit, coiffure et vêtements identiques à l'original."),
        (1, f"{p}bouche_X.png", "Bouche", "X — repos (silence)", f"Aucune génération : copiez {p}base.png et renommez la copie {p}bouche_X.png.", "Fichier identique à la base."),
        (2, f"{p}bouche_A.png", "Bouche", "A — lèvres fermées et pressées (M, B, P)", pre + "la bouche : lèvres fermées et légèrement pressées l'une contre l'autre, comme au début du son « M » ou « P », un peu plus fermes que dans la base ; yeux, sourcils et tout le reste inchangés." + post, "Lèvres jointes, différence subtile mais visible."),
        (3, f"{p}bouche_B.png", "Bouche", "B — dents serrées, bouche à peine ouverte (S, T, K, I)", pre + "la bouche : lèvres légèrement entrouvertes laissant voir les dents du haut et du bas serrées, ouverture très faible, pas de langue ; yeux et sourcils inchangés." + post, "Dents visibles et serrées, ouverture très faible."),
        (4, f"{p}bouche_C.png", "Bouche", "C — mi-ouverte (È, É)", pre + "la bouche : ouverte à moitié, forme légèrement ovale, dents du haut visibles, mâchoire un peu abaissée ; yeux et sourcils inchangés." + post, "Ouverture moyenne, dents du haut visibles."),
        (5, f"{p}bouche_D.png", "Bouche", "D — grande ouverte (A)", pre + "la bouche : grande ouverte, mâchoire nettement abaissée, forme ovale verticale, dents du haut visibles, intérieur visible ; yeux et sourcils inchangés." + post, "La plus grande ouverture de la série."),
        (6, f"{p}bouche_E.png", "Bouche", "E — arrondie ouverte (O)", pre + "la bouche : lèvres arrondies formant un « O » ouvert de taille moyenne, intérieur sombre visible ; yeux et sourcils inchangés." + post, "Cercle net, plus petit que D, plus grand que F."),
        (7, f"{p}bouche_F.png", "Bouche", "F — lèvres pincées et avancées (OU, U, W)", pre + "la bouche : lèvres pincées et projetées vers l'avant en un petit cercle serré, comme pour le son « OU » ; yeux et sourcils inchangés." + post, "Petit rond serré, lèvres avancées."),
        (8, f"{p}bouche_G.png", "Bouche", "G — lèvre inférieure sous les dents du haut (F, V)", pre + "la bouche : lèvre inférieure repliée vers l'intérieur, touchant les dents du haut visibles, lèvre supérieure normale, ouverture faible ; yeux et sourcils inchangés." + post, "Dents du haut posées sur la lèvre du bas."),
        (9, f"{p}bouche_H.png", "Bouche", "H — langue relevée (L)", pre + "la bouche : ouverte moyennement, pointe de la langue relevée et visible touchant l'arrière des dents du haut ; yeux et sourcils inchangés." + post, "Ouverture proche de C, langue visible."),
        (10, f"{p}yeux_fermes.png", "Yeux (clignement)", "Yeux fermés", pre + "les yeux : les deux paupières complètement fermées, chaque œil devient une ligne courbe douce (cils conservés s'il y en a), comme au milieu d'un clignement ; sourcils, bouche et tout le reste inchangés." + post, "Seules les paupières changent. Sourcils comme la base."),
        (11, f"{p}yeux_mi_clos.png", "Yeux (clignement)", "Yeux mi-clos", pre + "les yeux : paupières à moitié fermées, la moitié supérieure des yeux couverte, blanc et pupille encore visibles dessous, regard droit ; sourcils, bouche et tout le reste inchangés." + post, "Étape intermédiaire entre ouvert et fermé."),
        (12, f"{p}emotion_enjoue.png", "Émotion (sourcils + yeux)", "Enjoué", pre + "les sourcils et les yeux : sourcils légèrement relevés et détendus, yeux un peu plissés par un sourire (paupières inférieures remontées), expression joyeuse ; la bouche reste fermée et neutre comme dans la base." + post, "La bouche ne doit PAS changer."),
        (13, f"{p}emotion_surpris.png", "Émotion (sourcils + yeux)", "Surpris", pre + "les sourcils et les yeux : sourcils nettement relevés, yeux plus grands et bien ouverts, expression de surprise ; la bouche reste fermée et neutre comme dans la base." + post, "Sourcils hauts, yeux agrandis, bouche inchangée."),
        (14, f"{p}emotion_serieux.png", "Émotion (sourcils + yeux)", "Sérieux", pre + "les sourcils et les yeux : sourcils abaissés et légèrement froncés vers le centre, regard concentré et déterminé, sans colère ; la bouche reste fermée et neutre comme dans la base." + post, "Froncement léger, pas de colère, bouche inchangée."),
        (15, f"{p}emotion_pensif.png", "Émotion (sourcils + yeux)", "Pensif", pre + "les sourcils et les yeux : un sourcil légèrement relevé, l'autre au repos, regard tourné vers le haut et le côté comme quelqu'un qui réfléchit ; la bouche reste fermée et neutre comme dans la base." + post, "Regard décalé vers le haut, un seul sourcil relevé."),
        (16, f"{p}regard_gauche.png", "Regard (pupilles)", "Regard vers la gauche de l'image", pre + "les iris et pupilles des deux yeux sont déplacés vers la gauche de l'image (le personnage regarde vers sa droite, hors caméra) ; paupières, sourcils et bouche exactement comme la base." + post, "Seules les pupilles bougent."),
        (17, f"{p}regard_droite.png", "Regard (pupilles)", "Regard vers la droite de l'image", pre + "les iris et pupilles des deux yeux sont déplacés vers la droite de l'image (le personnage regarde vers sa gauche, hors caméra) ; paupières, sourcils et bouche exactement comme la base." + post, "Miroir de regard_gauche."),
        (18, f"{p}regard_haut.png", "Regard (pupilles)", "Regard vers le haut", pre + "les iris et pupilles des deux yeux sont déplacés vers le haut, comme quelqu'un qui cherche ses mots ; paupières ouvertes comme la base, sourcils et bouche inchangés." + post, "Pupilles en haut de l'œil."),
        (19, f"{p}bouche_X_sourire.png", "Bouche souriante", "X sourire — bouche fermée, coins relevés", pre + "la bouche : fermée avec un vrai sourire, coins des lèvres nettement relevés ; les yeux et les sourcils ne changent pas (pas de plissement)." + post, "Les yeux ne se plissent pas."),
        (20, f"{p}bouche_B_sourire.png", "Bouche souriante", "B sourire — grand sourire dents visibles", pre + "la bouche : grand sourire ouvert montrant les dents du haut serrées contre celles du bas, coins relevés ; yeux et sourcils inchangés." + post, "Dents visibles et serrées, ouverture faible."),
        (21, f"{p}bouche_D_sourire.png", "Bouche souriante", "D sourire — grande ouverture riante", pre + "la bouche : grande ouverte comme sur un « A » rieur, mâchoire abaissée, coins relevés, dents du haut visibles ; yeux et sourcils inchangés." + post, "Même ouverture que D avec les coins remontés."),
    ]
    if hands:
        r += [
            (22, f"{p}main_levee.png", "Mains (gestes)", "Salut — main droite levée, paume ouverte", pre + "le personnage lève sa main droite (à gauche de l'image) à hauteur de l'épaule, paume ouverte vers la caméra, comme pour dire bonjour ; main et avant-bras entièrement dans le cadre, sans cacher le visage ni le cou ; tête, visage, coiffure, épaules et fond inchangés." + post, "Main entière dans l'image, ne touche pas le visage."),
            (23, f"{p}mains_ouvertes.png", "Mains (gestes)", "Explication — deux mains ouvertes, paumes vers le haut", pre + "les deux mains apparaissent en bas de l'image à hauteur de la poitrine, paumes ouvertes vers le haut, coudes bas, comme quelqu'un qui explique calmement ; mains entièrement dans le cadre sous le menton ; tête, visage et fond inchangés." + post, "Deux mains visibles en bas, rien ne cache le visage."),
            (24, f"{p}main_index.png", "Mains (gestes)", "Point important — index levé", pre + "le personnage lève sa main droite (à gauche de l'image) à hauteur de l'épaule, index tendu vers le haut, autres doigts repliés ; main entièrement dans le cadre, sans cacher le visage ; tête, visage et fond inchangés." + post, "Même position de bras que main_levee."),
            (25, f"{p}main_pouce.png", "Mains (gestes)", "Approbation — pouce levé", pre + "le personnage montre un pouce levé de la main droite (à gauche de l'image), main à hauteur de l'épaule, tournée vers la caméra ; main entièrement dans le cadre, sans cacher le visage ; tête, visage et fond inchangés." + post, "Pouce lisible, main dans le cadre, visage inchangé."),
        ]
    return r


HEAD = ["Ordre", "Nom du fichier", "Calque", "Forme / état", "Prompt à coller dans ChatGPT (image jointe : l'original pour la ligne 0, la base ensuite)", "À vérifier sur le résultat"]
FILLS = {"Base": "F8CBAD", "Bouche": "FCE4D6", "Yeux (clignement)": "E2EFDA", "Émotion (sourcils + yeux)": "FFF2CC", "Regard (pupilles)": "E2F0D9", "Bouche souriante": "FBE5D6", "Mains (gestes)": "DDEBF7"}


def write_sheet(ws, rows, notes=None):
    hfill = PatternFill("solid", fgColor="1F4E78"); thin = Side(style="thin", color="BFBFBF"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for j, t in enumerate(HEAD, 1):
        c = ws.cell(row=1, column=j, value=t); c.font = Font(name="Arial", bold=True, color="FFFFFF", size=10); c.fill = hfill
        c.alignment = Alignment(vertical="center", wrap_text=True); c.border = border
    for i, r in enumerate(rows, 2):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v); c.font = Font(name="Arial", size=10, bold=(j == 2))
            c.alignment = Alignment(vertical="top", wrap_text=True); c.border = border
            if j <= 4: c.fill = PatternFill("solid", fgColor=FILLS[r[2]])
        ws.row_dimensions[i].height = 110
    for col, w in zip("ABCDEF", (7, 28, 24, 36, 95, 48)): ws.column_dimensions[col].width = w
    ws.row_dimensions[1].height = 32; ws.freeze_panes = "A2"
    if notes:
        start = len(rows) + 3
        ws.cell(row=start, column=2, value="Points d'attention propres à ce personnage").font = Font(name="Arial", bold=True, size=10)
        for k, n in enumerate(notes, 1):
            c = ws.cell(row=start + k, column=2, value="• " + n); c.font = Font(name="Arial", size=10); c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=start + k, start_column=2, end_row=start + k, end_column=6); ws.row_dimensions[start + k].height = 32


def write_guide(ws, prefixes):
    lines = [
        ("Nouveaux personnages : 26 images chacun pour la marionnette 2D", True), ("", False),
        ("Un onglet par personnage. Les noms de fichiers portent un préfixe (" + ", ".join(p + "_" for p in prefixes) + ") : déposez les images de chaque personnage dans assets/<préfixe>/ (" + ", ".join("assets/" + p + "/" for p in prefixes) + ").", False), ("", False),
        ("Étapes (par personnage)", True),
        ("1. Ouvrez UNE conversation ChatGPT par personnage. Joignez l'image originale et collez le prompt de la ligne 0 : la base doit avoir la bouche FERMÉE neutre et le regard droit. Vérifiez-la avec soin, tout en dépend.", False),
        ("2. Pour chaque ligne suivante, joignez la base (pas l'original) et collez le prompt. Demandez toujours une modification de l'image jointe, jamais une nouvelle image.", False),
        ("3. Enregistrez chaque résultat sous le nom exact de la colonne B, en PNG, même taille que la base.", False),
        ("4. Vérifiez l'alignement en alternant base et nouvelle image dans deux onglets : seule la zone annoncée doit bouger. Sinon : « Recommence en gardant exactement l'image jointe, ne modifie que … ».", False),
        ("5. Déposez les fichiers dans assets/<préfixe>/ et prévenez-moi : j'écris le manifeste, je vérifie alignement et détourage (et je recale une image légèrement décalée si besoin), le personnage apparaît dans Réglages → Personnage.", False), ("", False),
        ("Rappels valables pour tous", True),
        ("Les images d'émotion ne doivent pas changer la bouche ; les bouches ne doivent pas changer les yeux ni les sourcils. L'application tolère de petits écarts mais pas une bouche ouverte sur une image d'yeux.", False),
        ("Les mains sont optionnelles : sans elles, la piste gestes n'affiche simplement rien pour ce personnage. Elles doivent rester dans le cadre, sans toucher le visage.", False),
        ("Ordre conseillé : base, bouche_D, yeux_fermes, emotion_enjoue (premier test), puis le reste.", False),
    ]
    for i, (t, bold) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t); c.font = Font(name="Arial", size=11 if bold else 10, bold=bold); c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120


def main():
    out = sys.argv[1]
    args = sys.argv[2:]
    specs = []
    i = 0
    while i < len(args):
        prefix = args[i]; desc = args[i + 1]; i += 2
        notes = []; same = desc; hands = True
        while i < len(args) and args[i].startswith("--"):
            if args[i] == "--note": notes.append(args[i + 1]); i += 2
            elif args[i] == "--same": same = args[i + 1]; i += 2
            elif args[i] == "--sans-mains": hands = False; i += 1
            else: break
        specs.append((prefix, desc, same, hands, notes))
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    write_guide(wb.create_sheet("Mode d'emploi"), [s[0] for s in specs])
    for prefix, desc, same, hands, notes in specs:
        write_sheet(wb.create_sheet(prefix), rows_for(prefix, desc, same, hands), notes)
    wb.save(out); print("écrit :", out)


if __name__ == "__main__":
    main()
